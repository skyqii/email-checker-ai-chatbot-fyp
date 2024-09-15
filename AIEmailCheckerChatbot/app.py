from flask import Flask, request, jsonify, render_template
from werkzeug.utils import secure_filename
import os
import re
import PyPDF2  # For PDF handling
import docx  # For Word document handling
import openpyxl  # For Excel handling

app = Flask(__name__)

UPLOAD_FOLDER = 'uploads/'
ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx', 'xlsx', 'eml', 'msg'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# check allowed file types
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(filepath, file_type):
    content = ""
    if file_type == 'txt':
        with open(filepath, 'r') as file:
            content = file.read().lower()
    elif file_type == 'pdf':
        with open(filepath, 'rb') as file:
            reader = PyPDF2.PdfFileReader(file)
            content = ''.join([reader.getPage(i).extract_text() for i in range(reader.numPages)]).lower()
    elif file_type == 'docx':
        doc = docx.Document(filepath)
        content = ' '.join([para.text for para in doc.paragraphs]).lower()
    elif file_type == 'xlsx':
        workbook = openpyxl.load_workbook(filepath)
        sheet = workbook.active
        content = ' '.join([str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]).lower()
    return content

def classify_file(filepath):

    file_type = filepath.rsplit('.', 1)[1].lower()
    content = extract_text_from_file(filepath, file_type)
    
    # classif logic here

# Route for file uploads
@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        classification = classify_file(filepath)
        return jsonify({'classification': classification}), 200
    
    return jsonify({'error': 'File type not allowed'}), 400

@app.route('/')
def index():
    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True)
